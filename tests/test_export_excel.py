"""Tests for export_excel.py — the 500-day workbook from the same store."""

from datetime import datetime, timedelta, timezone

import pandas as pd
import pytest
from freezegun import freeze_time
from openpyxl import load_workbook

import analytics
import export_excel
import integrity
import journal
import store as store_mod

NOW = datetime(2026, 6, 28, 12, 0, 0, tzinfo=timezone.utc)
CORE = [50, 60, 70, 100, 140, 150, 180, 200, 250, 300]
EXPECTED_SHEETS = ["Cover", "Glucose_5min", "Daily_Metrics", "Meals", "Food_Rank",
                   "Experiments", "Wearables", "Supplements", "Symptoms_Mood", "Labs",
                   "AGP_Profile"]


@pytest.fixture()
def st():
    with freeze_time(NOW):
        integrity.cache_clear()
        s = store_mod.Store(":memory:")
        s.append_glucose(pd.DataFrame(
            [{"measured_at": NOW - timedelta(minutes=i), "glucose_mgdl": v}
             for i, v in enumerate(CORE)]), now=NOW)
        # a supplement + a mood/symptom log + a lab marker
        raws = [{"entry_id": "s1", "date": "2026-06-28", "time": "08:00", "type": "supplement",
                 "item": "Triphala", "tags": "supplement"},
                {"entry_id": "m1", "date": "2026-06-28", "time": "21:00", "type": "mood",
                 "mood_1to5": "3", "energy_1to5": "2", "symptom": "fatigue",
                 "symptom_sev_1to5": "2", "tags": ""}]
        s.append_events("log", journal.parse_rows(raws), key_field="key", now=NOW)
        s.append_events("labs", [{"key": "crp", "measured_at": (NOW - timedelta(days=2)).isoformat(),
                                  "marker": "CRP", "value": 5.0, "unit": "mg/L", "ref_high": 3.0}],
                        key_field="key", now=NOW)
        yield s
        s.close()
        integrity.cache_clear()


def test_workbook_has_all_sheets(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
    assert wb.sheetnames == EXPECTED_SHEETS


def test_glucose_rowcount_matches_store(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
    assert wb["Glucose_5min"].max_row - 1 == len(st.glucose_all()) == 10
    assert export_excel.self_check(st, wb) == []


def test_glucose_conditional_formatting_present(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
    rules = list(wb["Glucose_5min"].conditional_formatting)
    assert rules, "expected glucose conditional formatting rules"


def test_daily_metric_matches_dashboard(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
        trend = analytics.compute(st, window_days=500, now=NOW, group_by="day")
    # spot-check: the Daily_Metrics row for a day equals the analytics (dashboard) value
    series_by_day = {r["group"]: r for r in trend["series"]}
    ws = wb["Daily_Metrics"]
    headers = [c.value for c in ws[1]]
    mean_col = headers.index("mean_mgdl")
    for row in ws.iter_rows(min_row=2, values_only=True):
        day = row[0]
        assert row[mean_col] == pytest.approx(series_by_day[day]["mean_mgdl"])


def test_cover_has_integrity_statement(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
    text = "\n".join(str(c.value) for r in wb["Cover"].iter_rows() for c in r if c.value)
    assert "computed deterministically" in text
    assert "Date range" in text


def test_labs_high_marker_flagged(st):
    with freeze_time(NOW):
        wb = export_excel.build_workbook(st, now=NOW)
    assert wb["Labs"].max_row - 1 == 1
    assert list(wb["Labs"].conditional_formatting), "expected lab high-marker CF"


def test_write_and_reload(st, tmp_path):
    with freeze_time(NOW):
        path = export_excel.write_excel(st, str(tmp_path / "HealthOS_500d.xlsx"), now=NOW)
    wb = load_workbook(path)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb["Glucose_5min"]["B2"].value in [float(v) for v in CORE]
