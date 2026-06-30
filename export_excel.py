"""export_excel.py — the 500-day workbook (§D), generated ON DEMAND from the SAME store.

HealthOS_500d.xlsx with sheets: Cover, Glucose_5min, Daily_Metrics, Meals, Food_Rank,
Experiments, Wearables, Supplements, Symptoms_Mood, Labs, AGP_Profile. Frozen + bold headers,
number formats, column auto-fit, conditional formatting for glucose (red >180 / amber 140-180 /
green 70-140 / dark-red <70) and high lab markers, and a Cover sheet with the date range,
per-source last-sync ages, and the integrity statement.

Because it reads the SAME store the dashboard reads (§A rule 7), the workbook and the dashboard
cannot disagree. Triggers: the dashboard "Export" button (links to the generated file), the
``healthhub-export`` CLI, and the weekly GitHub Action.
"""

from __future__ import annotations

from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analytics
import clinical
import experiments as experiments_mod
import food_impact
import heartbeat as heartbeat_mod
import integrity
import journal

LOCAL_TZ = "Asia/Dubai"
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
INTEGRITY_STATEMENT = (
    "Every number in this workbook is computed deterministically by Python from the same "
    "500-day store the live dashboard reads — never by an LLM, never interpolated. "
    "Out-of-range or future rows are quarantined, not plotted. Values carry their source and "
    "freshness; gaps are shown as gaps. Behaviour/analytics tooling only — not medical advice.")


def _local_naive(ts_iso):
    """UTC ISO -> naive Asia/Dubai datetime (openpyxl can't store tz-aware datetimes)."""
    dt = integrity.freshness(ts_iso, now=integrity.now_utc()).measured_at
    return dt.astimezone(ZoneInfo(LOCAL_TZ)).replace(tzinfo=None)


def _autofit(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(42, max(10, width + 2))


def _sheet(wb, title, headers, rows, *, number_formats=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = _HEADER_FONT, _HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if number_formats:
        for col_idx, fmt in number_formats.items():
            for c in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                c[0].number_format = fmt
    _autofit(ws)
    return ws


def _glucose_cf(ws, col_letter, nrows):
    if nrows < 1:
        return
    rng = f"{col_letter}2:{col_letter}{nrows + 1}"
    fills = {"red": "F8696B", "amber": "FFEB84", "green": "63BE7B", "low": "C00000"}
    add = ws.conditional_formatting.add
    add(rng, CellIsRule(operator="greaterThan", formula=["180"], stopIfTrue=True,
                        fill=PatternFill("solid", fgColor=fills["red"])))
    add(rng, CellIsRule(operator="lessThan", formula=["70"], stopIfTrue=True,
                        fill=PatternFill("solid", fgColor=fills["low"]),
                        font=Font(color="FFFFFF")))
    add(rng, CellIsRule(operator="between", formula=["140", "180"], stopIfTrue=True,
                        fill=PatternFill("solid", fgColor=fills["amber"])))
    add(rng, CellIsRule(operator="between", formula=["70", "140"], stopIfTrue=True,
                        fill=PatternFill("solid", fgColor=fills["green"])))


def build_workbook(store, *, now=None, wearables=None, health=None, window_days: int = 500):
    now = integrity.now_utc() if now is None else now
    wb = Workbook()

    # --- gather (all from the SAME store) ---
    gdf = store.glucose_all()
    trend = analytics.compute(store, window_days=window_days, now=now, group_by="day")
    logs = journal.records(store)
    meals = journal.meals(store)

    # ---- Cover (first sheet) ----
    if len(gdf):
        lo = _local_naive(gdf["measured_at"].min().isoformat())
        hi = _local_naive(gdf["measured_at"].max().isoformat())
        date_range = f"{lo.date()} → {hi.date()}"
    else:
        date_range = "no data"
    hb = health or heartbeat_mod.status(
        {"glucose": (store.latest_glucose() or {}).get("ts")}, now=now)
    cover_rows = [["Date range", date_range], ["Generated", now.isoformat()],
                  ["Glucose readings", len(gdf)], ["Days", trend["n_readings"] and len(trend["series"])],
                  ["", ""], ["Source", "Last sync / state"]]
    for s in hb.get("sources", []):
        cover_rows.append([s["source"], f"{s.get('age') or 'no data'} ({s['state']})"])
    cover_rows += [["", ""], ["Integrity", INTEGRITY_STATEMENT]]
    cover = _sheet(wb, "Cover", ["HealthOS 500-day export", ""], cover_rows, first=True)
    cover["A1"].font = Font(bold=True, size=14)

    # ---- Glucose_5min ----
    grows = [[_local_naive(r.measured_at.isoformat()), float(r.glucose_mgdl), r.source]
             for r in gdf.itertuples()]
    gws = _sheet(wb, "Glucose_5min", ["timestamp", "glucose_mgdl", "source"], grows,
                 number_formats={1: "yyyy-mm-dd hh:mm", 2: "0"})
    _glucose_cf(gws, "B", len(grows))

    # ---- Daily_Metrics ----
    cols = ["mean_mgdl", "gmi_pct", "cv_pct", "tir_pct", "titr_pct", "tbr_pct", "tar_pct",
            "gri", "mage_mgdl"]
    drows = [[row["group"], row["n"]] + [row.get(c) for c in cols] for row in trend["series"]]
    _sheet(wb, "Daily_Metrics", ["date", "n"] + cols, drows,
           number_formats={i: "0.0" for i in range(3, 3 + len(cols))})

    # ---- Meals ----
    mrows = []
    for m in meals:
        r = food_impact.meal_response(store, m, now=now)
        mrows.append([_local_naive(r["measured_at"]), m.get("item"), m.get("net_carbs_g"),
                      r.get("baseline_mgdl"), r.get("delta_peak_mgdl"), r.get("iauc_120"),
                      r.get("time_to_peak_min"), r.get("per_gram"), "; ".join(m.get("tags") or [])])
    _sheet(wb, "Meals", ["timestamp", "item", "net_carbs_g", "baseline", "delta_peak",
                         "iauc_120", "time_to_peak_min", "per_gram", "tags"], mrows,
           number_formats={1: "yyyy-mm-dd hh:mm"})

    # ---- Food_Rank ----
    frows = [[f["item"], f["n"], f["mean_delta_peak_mgdl"], f["mean_iauc_120"],
              f.get("mean_per_gram")] for f in food_impact.rank_foods(store, now=now)]
    _sheet(wb, "Food_Rank", ["item", "n", "mean_delta_peak", "mean_iauc_120", "mean_per_gram"],
           frows, number_formats={3: "0.0", 4: "0.0", 5: "0.00"})

    # ---- Experiments ----
    erows = [[e["tag"], e.get("effect_abs"), e.get("effect_pct"), e.get("n_treated"),
              e.get("n_control"), e.get("signal_strength"), e.get("causal_label")]
             for e in experiments_mod.compare_all_tags(store, now=now)]
    _sheet(wb, "Experiments", ["tag", "effect_abs", "effect_pct", "n_treated", "n_control",
                              "signal", "causal_label"], erows)

    # ---- Wearables ----
    wrows = []
    for date, d in sorted((wearables or {}).get("days", {}).items()):
        wrows.append([date, d.get("steps"), d.get("total_calories"), d.get("active_calories"),
                      d.get("weight_kg"), d.get("hydration_ml"), d.get("sleep_total_min"),
                      d.get("hr_avg"), d.get("spo2_avg")])
    _sheet(wb, "Wearables", ["date", "steps", "total_cal", "active_cal", "weight_kg",
                            "hydration_ml", "sleep_min", "hr_avg", "spo2_avg"], wrows)

    # ---- Supplements ----
    srows = [[_local_naive(r.get("measured_at") or r.get("ts")), r.get("item") or r.get("supplements"),
              "; ".join(r.get("tags") or [])]
             for r in logs if r.get("type") == "supplement" or r.get("supplements")]
    _sheet(wb, "Supplements", ["timestamp", "supplement", "tags"], srows,
           number_formats={1: "yyyy-mm-dd hh:mm"})

    # ---- Symptoms_Mood ----
    smrows = [[_local_naive(r.get("measured_at") or r.get("ts")), r.get("symptom"),
               r.get("symptom_sev_1to5"), r.get("mood_1to5"), r.get("energy_1to5")]
              for r in logs if r.get("symptom") or r.get("mood_1to5") is not None
              or r.get("energy_1to5") is not None]
    _sheet(wb, "Symptoms_Mood", ["timestamp", "symptom", "severity", "mood", "energy"], smrows,
           number_formats={1: "yyyy-mm-dd hh:mm"})

    # ---- Labs (events stream 'labs'; high markers flagged) ----
    lrows = []
    for ev in store.events("labs"):
        lrows.append([_local_naive(ev.get("measured_at") or ev.get("ts")), ev.get("marker"),
                      ev.get("value"), ev.get("unit"), ev.get("ref_high")])
    lws = _sheet(wb, "Labs", ["timestamp", "marker", "value", "unit", "ref_high"], lrows,
                 number_formats={1: "yyyy-mm-dd hh:mm"})
    if lrows:  # flag a value above its reference high (per-row formula CF)
        lws.conditional_formatting.add(
            f"C2:C{len(lrows) + 1}",
            CellIsRule(operator="greaterThan", formula=["E2"],
                       fill=PatternFill("solid", fgColor="F8696B")))

    # ---- AGP_Profile ----
    arows = []
    for slot, band in enumerate(trend.get("agp", [])):
        hhmm = f"{slot * 15 // 60:02d}:{slot * 15 % 60:02d}"
        if band is None:
            arows.append([hhmm, None, None, None, None, None])
        else:
            arows.append([hhmm, band["p10"], band["p25"], band["p50"], band["p75"], band["p90"]])
    _sheet(wb, "AGP_Profile", ["slot", "p10", "p25", "p50", "p75", "p90"], arows,
           number_formats={i: "0" for i in range(2, 7)})
    return wb


def write_excel(store, path: str = "HealthOS_500d.xlsx", *, now=None, wearables=None,
                health=None, window_days: int = 500) -> str:
    wb = build_workbook(store, now=now, wearables=wearables, health=health,
                        window_days=window_days)
    wb.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    import argparse

    import glucose as glucose_mod
    import store as store_mod

    p = argparse.ArgumentParser(prog="healthhub-export", description=__doc__)
    p.add_argument("--db", default="healthhub.db", help="store path (sqlite)")
    p.add_argument("--csv", default=None, help="optional glucose CSV to ingest before export")
    p.add_argument("-o", "--output", default="HealthOS_500d.xlsx")
    args = p.parse_args(argv)

    st = store_mod.Store(args.db)
    if args.csv:
        import pandas as pd
        df = pd.read_csv(args.csv)
        if "measured_at" in df.columns:
            df["measured_at"] = pd.to_datetime(df["measured_at"], utc=False)
        st.append_glucose(df, source="csv")
    path = write_excel(st, args.output)
    print(f"wrote {path}")
    return 0


def self_check(store, wb) -> list[str]:
    """Invariant: the Glucose_5min sheet row count equals the store's glucose row count."""
    violations = []
    n_store = len(store.glucose_all())
    n_sheet = wb["Glucose_5min"].max_row - 1 if "Glucose_5min" in wb.sheetnames else -1
    if n_sheet != n_store:
        violations.append(f"Glucose_5min rows {n_sheet} != store {n_store}")
    return violations


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
